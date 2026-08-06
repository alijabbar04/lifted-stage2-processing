"""Persistent misnaming record for the Stage 2 pipeline.

One xlsx, appended to whenever a misnamed document is identified (by the
user, by a Claude review pass, or by any future check). Lives with Stage 2's
other data files in %APPDATA%\\DocReviewAIStation.

Columns:
  date_identified | care_home | worker | file | wrong_name | correct_name |
  found_by | status ('Resolved' / 'Open - needs human check') |
  date_resolved | notes
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

LOG_PATH = (Path.home() / "AppData" / "Roaming" / "DocReviewAIStation"
            / "Misnaming Record.xlsx")
COLUMNS = ["date_identified", "care_home", "worker", "file", "wrong_name",
           "correct_name", "found_by", "status", "date_resolved", "notes"]


def _open():
    if LOG_PATH.exists():
        wb = load_workbook(LOG_PATH)
        return wb, wb.active
    wb = Workbook()
    ws = wb.active
    ws.title = "Misnames"
    ws.append(COLUMNS)
    for col, width in zip("ABCDEFGHIJ", (12, 22, 18, 42, 34, 34, 26, 24, 12, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    return wb, ws


def add_rows(rows):
    """rows: list of dicts keyed by COLUMNS (missing keys become '')."""
    wb, ws = _open()
    for r in rows:
        ws.append([str(r.get(c, "")) for c in COLUMNS])
    wb.save(LOG_PATH)
    return LOG_PATH


def add(care_home, worker, file, wrong_name, correct_name, found_by,
        status="Resolved", notes="", date_identified=None, date_resolved=None):
    today = date.today().isoformat()
    return add_rows([{
        "date_identified": date_identified or today,
        "care_home": care_home, "worker": worker, "file": file,
        "wrong_name": wrong_name, "correct_name": correct_name,
        "found_by": found_by, "status": status,
        "date_resolved": (date_resolved or (today if status == "Resolved" else "")),
        "notes": notes,
    }])
