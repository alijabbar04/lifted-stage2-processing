"""Shared pipeline contract, v1. Source of truth: Stage 1/src.

Vendored unchanged into each standalone app by tools/sync_pipeline_shared.py.
No UI/application imports; tests can exercise handover without a browser or API.
"""
from __future__ import annotations

import contextlib
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import time
import unicodedata
import uuid
from collections import Counter
from difflib import SequenceMatcher

CONTRACT_VERSION = "1"
BUILD_VERSION = "2026.09.04-rc1"
FIELDS = ["schema_version", "worker_key", "care_home", "agency_name",
          "folder_name", "original_folder_name", "source_path", "file_count",
          "duplicate_flag", "status", "matched_name", "matched_id", "score",
          "approved", "duplicate", "top_candidates", "profile_url",
          "stage1_state", "stage2_state", "stage3_state", "updated_at"]
MATCH_FIELDS = ["status", "matched_name", "matched_id", "score", "approved",
                "duplicate", "top_candidates", "profile_url", "agency_name"]
DOCUMENT_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff",
                       ".doc", ".docx", ".heic", ".webp", ".bmp", ".rtf"}
STATE_SUFFIX = re.compile(r"\s*(?:\[(Files|Processed|Uploaded|Approved)\]|"
                          r"(Files|Processed|Uploaded|Approved))\s*$", re.I)


def normalise(value):
    value = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", str(value or ""))
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return " ".join(re.findall(r"[a-z0-9]+", value.lower()))


def flag(value):
    return str(value or "").strip().lower() in {"yes", "y", "true", "1"}


def profile_id(value):
    value = str(value or "").strip()
    if re.fullmatch(r"\d+\.0", value):
        value = value[:-2]
    return value if value.isdigit() and int(value) > 0 else ""


def normalise_row(raw):
    row = {str(k).strip(): "" if v is None else str(v).strip()
           for k, v in raw.items() if k is not None}
    if not row.get("score") and "match_score" in row:
        row["score"] = row.get("match_score", "")
    row.pop("match_score", None)
    if "matched_id" in row:
        row["matched_id"] = profile_id(row["matched_id"])
    return row


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            # A historical workbook may open on its Summary sheet.
            for sheet in wb:
                values = iter(sheet.iter_rows(values_only=True))
                headers = [str(v or "").strip() for v in next(values, ())]
                if "folder_name" in headers:
                    return [normalise_row(dict(zip(headers, cells))) for cells in values
                            if any(v is not None for v in cells)]
            raise ValueError("Excel file has no worker roster sheet (folder_name column).")
        finally:
            wb.close()
    if path.suffix.lower() == ".xls":
        raise ValueError("Please save this legacy .xls file as .xlsx or CSV first.")
    with path.open(newline="", encoding="utf-8-sig") as stream:
        reader = csv.DictReader(stream)
        if "folder_name" not in (reader.fieldnames or []):
            raise ValueError("This is not a worker roster: folder_name column is missing.")
        return [normalise_row(row) for row in reader if any(row.values())]


def atomic_csv(path, rows, fields=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(fields or FIELDS)
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    handle, name = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(handle, "w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="raise")
            writer.writeheader()
            writer.writerows(rows)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)
    return path


@contextlib.contextmanager
def roster_lock(path, timeout=10):
    """Cross-process merge lock: stages may legitimately run concurrently."""
    lock = Path(str(path) + ".lock")
    lock.parent.mkdir(parents=True, exist_ok=True)
    stream = lock.open("a+b")
    stream.seek(0, 2)
    if stream.tell() == 0:
        stream.write(b"0")
        stream.flush()
    deadline = time.monotonic() + timeout
    acquired = False
    try:
        while not acquired:
            try:
                stream.seek(0)
                if os.name == "nt":
                    import msvcrt
                    msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl
                    fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
                acquired = True
            except OSError:
                if time.monotonic() >= deadline:
                    raise TimeoutError("Roster is being updated by another stage; try again.")
                time.sleep(.05)
        yield
    finally:
        if acquired:
            stream.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(stream, fcntl.LOCK_UN)
        stream.close()


def care_name(source):
    name = Path(source).name
    return STATE_SUFFIX.sub("", name).strip() or name


def care_root(source):
    source = Path(source)
    if STATE_SUFFIX.search(source.name):
        # Standard nested layout. Legacy Desktop/<Care> Approved is a sibling.
        if normalise(source.parent.name) == normalise(care_name(source)):
            return source.parent
        sibling = source.parent / care_name(source)
        if sibling.is_dir():
            return sibling
    return source


def roster_path(source, care_home=None):
    root = care_root(source)
    name = re.sub(r'[<>:"/\\|?*]', "_", care_home or care_name(source)).strip()
    return root / f"{name} roster.csv"


def find_roster(source, care_home=None, extra_dirs=()):
    source = Path(source)
    care = care_home or care_name(source)
    roots = list(dict.fromkeys([care_root(source), source, *map(Path, extra_dirs)]))
    exact_names = [f"{care} roster", f"{care}_match_review",
                   f"{care} match review", f"{care} name match", f"{care}_name_match"]
    stems = {normalise(n) for n in exact_names}
    canonical = roster_path(source, care)
    if canonical.is_file():
        return canonical
    for root in roots:
        if not root.is_dir():
            continue
        for item in sorted(root.iterdir()):
            if item.suffix.lower() in {".csv", ".xlsx", ".xlsm"} and normalise(item.stem) in stems:
                return item
        # Generic legacy roster is only valid within the explicitly chosen care-home tree.
        legacy = root / "roster.csv"
        if root in {source, care_root(source)} and legacy.is_file():
            return legacy
    return None


def merge_rows(existing, updates, owned_fields=None, care_home="", root=None):
    result = [normalise_row(r) for r in existing]
    for update in updates:
        update = normalise_row(update)
        key, folder = update.get("worker_key"), update.get("folder_name", "").casefold()
        candidates = [i for i, r in enumerate(result)
                      if (key and key == r.get("worker_key")) or
                      (folder and folder == r.get("folder_name", "").casefold())]
        if len(candidates) > 1:
            raise ValueError(f"Duplicate roster rows for {update.get('folder_name')}; review before saving.")
        row = result[candidates[0]] if candidates else {}
        for k, v in update.items():
            if owned_fields is None or k in owned_fields or not row:
                row[k] = v
        row.setdefault("folder_name", update.get("folder_name", ""))
        row.setdefault("original_folder_name", row["folder_name"])
        row.setdefault("care_home", care_home)
        if not row.get("worker_key"):
            identity = str(root or care_home).casefold() + "/" + row["original_folder_name"].casefold()
            row["worker_key"] = str(uuid.uuid5(uuid.NAMESPACE_URL, identity))
        row["schema_version"] = CONTRACT_VERSION
        row["updated_at"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
        if not candidates:
            result.append(row)
    return result


def update_roster(path, updates, owned_fields=None, care_home="", seed=()):
    path = Path(path)
    with roster_lock(path):
        existing = read_rows(path) if path.exists() else list(seed)
        rows = merge_rows(existing, updates, owned_fields, care_home, path.parent)
        atomic_csv(path, rows)
    return rows


def worker_dirs(source):
    source = Path(source)
    excluded = {"bulk", "overwrite documents", "other", "archived reports"}
    return [p for p in sorted(source.iterdir()) if p.is_dir()
            and not p.name.startswith((".", "_"))
            and normalise(p.name) not in excluded
            and not STATE_SUFFIX.search(p.name)
            and "report" not in p.name.lower()]


def scan_workers(source, care_home=None):
    source = Path(source)
    root = care_root(source)
    rows = []
    for folder in worker_dirs(source):
        files = [p for p in folder.rglob("*") if p.is_file()
                 and p.suffix.lower() in DOCUMENT_EXTENSIONS]
        rows.append({"folder_name": folder.name, "source_path": os.path.relpath(folder, root),
                     "file_count": str(len(files)), "care_home": care_home or care_name(source)})
    return rows


def refresh_roster(source, care_home=None, stage=None):
    path = roster_path(source, care_home)
    old = find_roster(source, care_home)
    seed = read_rows(old) if old and old != path else []
    rows = scan_workers(source, care_home)
    if stage:
        for row in rows:
            row[f"{stage}_state"] = "discovered"
    update_roster(path, rows, care_home=care_home or care_name(source), seed=seed)
    return path


def record_worker_move(source, destination, stage, state="complete"):
    source, destination = Path(source), Path(destination)
    path = roster_path(source.parent)
    old = find_roster(source.parent)
    if not old:
        return None
    with roster_lock(path):
        rows = read_rows(path if path.exists() else old)
        matches = [r for r in rows if r.get("folder_name", "").casefold() == source.name.casefold()]
        if len(matches) != 1:
            raise ValueError(f"Expected one roster row for {source.name}; found {len(matches)}.")
        if any(r is not matches[0] and r.get("folder_name", "").casefold() == destination.name.casefold() for r in rows):
            raise ValueError(f"Destination name {destination.name} already belongs to another roster worker.")
        matches[0].update(folder_name=destination.name,
            source_path=os.path.relpath(destination, path.parent),
            file_count=str(sum(p.is_file() and p.suffix.lower() in DOCUMENT_EXTENSIONS
                               for p in destination.rglob("*"))),
            **{f"{stage}_state": state, "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z")})
        atomic_csv(path, rows)
    return path


def name_score(folder_name, platform_name):
    f, p = set(normalise(folder_name).split()), set(normalise(platform_name).split())
    if not f or not p:
        return 0.0
    return .7 * len(f & p) / len(f) + .3 * SequenceMatcher(
        None, normalise(folder_name), normalise(platform_name)).ratio()


def classify(folder_name, roster):
    ranked = sorted(((name_score(folder_name, r["name"]), r) for r in roster),
                    key=lambda x: x[0], reverse=True)
    if not ranked:
        return "NONE", None, [], 0.0
    score, candidate = ranked[0]
    second = ranked[1][0] if len(ranked) > 1 else 0
    status = "CONFIDENT" if score >= .85 and score - second >= .15 else "FUZZY" if score >= .45 else "NONE"
    return status, candidate if status != "NONE" else None, ranked, score


def match_folders(folder_names, roster):
    counts = Counter(normalise(r["name"]) for r in roster)
    rows = []
    for name in folder_names:
        status, candidate, ranked, score = classify(name, roster)
        wid = profile_id(candidate["id"]) if candidate else ""
        duplicate = bool(candidate and counts[normalise(candidate["name"])] > 1)
        rows.append({"folder_name": name, "status": "DUPLICATE" if duplicate else status,
                     "matched_id": wid, "matched_name": candidate["name"] if candidate else "",
                     "score": f"{score:.2f}", "approved": "yes" if status == "CONFIDENT" and not duplicate else "",
                     "duplicate": "yes" if duplicate else "",
                     "top_candidates": " | ".join(f"{r['name']} ({r['id']}) {s:.2f}" for s, r in ranked[:3]),
                     "profile_url": f"https://app.lifted-talent.com/applicant/profile/{wid}/documents" if wid else ""})
    # One profile cannot be auto-approved for two different local workers.
    ids = Counter(r["matched_id"] for r in rows if r["matched_id"])
    for row in rows:
        if row["matched_id"] and ids[row["matched_id"]] > 1:
            row.update(status="DUPLICATE", duplicate="yes", approved="")
    return rows


_GENERIC_AGENCY = {"care", "home", "homes", "limited", "ltd", "agency", "services", "service", "the", "and"}


def agency_tokens(value):
    return set(normalise(value).split()) - _GENERIC_AGENCY


def agency_label_matches(query, label):
    tokens = agency_tokens(query)
    return bool(tokens and tokens <= agency_tokens(label) and normalise(label) != "all")


def agency_searches(query):
    # Full label first, then broad discovery using each distinctive token.
    tokens = [t for t in normalise(query).split() if t not in _GENERIC_AGENCY]
    return list(dict.fromkeys([query, " ".join(tokens), *reversed(tokens)]))


def choose_agency(query, labels):
    labels = list(dict.fromkeys(label.strip() for label in labels if label.strip()))
    exact = [label for label in labels if normalise(label) == normalise(query)]
    if len(exact) == 1:
        return exact[0]
    matches = [label for label in labels if agency_label_matches(query, label)]
    return matches[0] if len(matches) == 1 else None


def selected_agency(page):
    controls = page.locator('div[tabindex="0"]').filter(has=page.locator('span.truncate'))
    # Confirmed Lifted list layout: Agency is the first such control.
    if controls.count():
        return controls.first.locator('span.truncate').first.inner_text(timeout=1500).strip()
    return ""


def select_agency(page, query, log=lambda *_: None):
    """Bounded automatic discovery; never reads an unfiltered applicant list.

    Returns the verified real label, or an empty string for manual fallback.
    Browser adapter kept here so Stages 1 and 3 cannot drift.
    """
    try:
        current = selected_agency(page)
        if normalise(current) == normalise(query):
            return current
        control = page.locator('div[tabindex="0"]').filter(has=page.locator('span.truncate')).first
        control.click(timeout=4000)
        box = page.locator('input[type="text"]:visible').last
        labels = set()
        for search in agency_searches(query):
            box.fill(search, timeout=3000)
            page.wait_for_timeout(500)
            options = page.locator('[data-type="option"]:visible, [role="option"]:visible')
            found = options.all_text_contents()
            if not found:
                # Lifted's custom selector has no role=option. Read only its
                # dropdown container, never arbitrary page text/profile names.
                found = box.evaluate("el => (el.parentElement.parentElement.innerText || '').split('\\n')")
            labels.update(s.strip() for s in found if s.strip() and normalise(s) not in {"all", "no options", "agency"})
        chosen = choose_agency(query, sorted(labels))
        if not chosen:
            box.fill("", timeout=2000)
            box.press("Escape")
            log(f"Agency needs selection: {query}; candidates: {', '.join(sorted(labels)) or 'none'}")
            return ""
        box.fill(chosen, timeout=3000)
        page.wait_for_timeout(400)
        option = page.get_by_text(chosen, exact=True)
        if option.count() == 1:
            option.click(timeout=3000)
        else:
            box.press("Enter")
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            actual = selected_agency(page)
            if normalise(actual) == normalise(chosen):
                log(f"Agency selected: {actual}")
                return actual
            page.wait_for_timeout(200)
    except Exception as exc:
        log(f"Automatic agency selection needs help: {exc}")
    return ""


def settings_path(stage):
    root = Path(os.environ.get("LOCALAPPDATA") or (Path.home() / "AppData/Local"))
    return root / "Lifted" / "Settings" / f"{stage}.json"


def load_settings(stage):
    try:
        return json.loads(settings_path(stage).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_settings(stage, **updates):
    path = settings_path(stage)
    values = load_settings(stage)
    values.update(updates)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(values, stream, indent=2)
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def build_identity(stage):
    import sys
    if getattr(sys, "frozen", False):
        hasher = hashlib.sha256()
        with Path(sys.executable).open("rb") as stream:
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                hasher.update(block)
        return f"{stage} {BUILD_VERSION} | roster v{CONTRACT_VERSION} | executable {hasher.hexdigest()[:16]}"
    source = Path(__file__)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()[:12] if source.exists() else "bundled"
    return f"{stage} {BUILD_VERSION} | roster v{CONTRACT_VERSION} | shared {digest}"
