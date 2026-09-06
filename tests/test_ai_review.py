"""Offline repair contracts: complete peers, source ranking, bytes and records."""
import csv
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import ai_review as review


class Fixture:
    def __init__(self, base):
        self.base = Path(base)
        self.documents = self.base / "Processed"
        self.documents.mkdir()
        self.source = self.base / "source"
        (self.source / "src").mkdir(parents=True)
        shutil.copy2(Path(__file__).resolve().parents[1] / "src" / "Stage2_Processing.pyw",
                     self.source / "src" / "Stage2_Processing.pyw")
        self.request = self.base / "request"
        self.request.mkdir()

    def file(self, rel, data=None):
        path = self.documents / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data or ("synthetic " + rel).encode())
        return path

    def prepare(self, candidate, confidence=95, suffix="csv"):
        self.audit = self.base / f"audit.{suffix}"
        headers = ["Current Filename", "Suggested Filename", "Full File Path", "Confidence Score", "Review Status"]
        row = [candidate.name, "DBS Document.pdf", str(candidate), confidence, "Likely Misnamed"]
        if suffix == "csv":
            with self.audit.open("w", newline="", encoding="utf-8-sig") as stream:
                writer = csv.writer(stream)
                writer.writerow(headers)
                writer.writerow(row)
        else:
            from openpyxl import Workbook
            book = Workbook()
            book.active.title = "Audit"
            book.active.append(headers)
            book.active.append(row)
            book.save(self.audit)
            book.close()
        self.queue_path = self.request / "review_queue.json"
        self.queue = review.prepare(self.audit, "Synthetic Home", self.documents,
                                   self.source, self.queue_path)
        return self.queue

    def decisions(self, kind="DBS Document", scores=None):
        scores = scores or {}
        decisions = {"run_id": self.queue["run_id"], "reviewer": "Offline reviewer",
            "decisions": [{"candidate_id": self.queue["candidates"][0]["candidate_id"],
                           "decision": "rename", "approved_type": kind,
                           "review_notes": "Synthetic test decision", "evidence_summary": "Synthetic page evidence",
                           "pages_examined": [1]}], "peer_reviews": []}
        for row in self.queue["inventory"]:
            decisions["peer_reviews"].append({"path": row["path"], "sha256": row["sha256"],
                "score": scores.get(row["path"], 50), "legible": True, "complete": True,
                "signed": False, "date": "2025-01-01", "evidence_summary": "Synthetic quality evidence",
                "pages_examined": [1]})
        return decisions

    def plan(self, decisions):
        self.decisions_path = self.request / "decisions.json"
        review.write_json(self.decisions_path, decisions)
        self.plan_path = self.request / "apply_plan.json"
        return review.make_plan(self.queue_path, self.decisions_path, self.plan_path)


class TestAIReview(unittest.TestCase):
    def ranking_fixture(self, directory):
        f = Fixture(directory)
        existing = f.file("Worker/Overwrite Documents/DBS Document.pdf")
        incoming = f.file("Worker/Bulk/Batch 01/Other - Unknown.pdf")
        f.prepare(incoming)
        decisions = f.decisions(scores={existing.relative_to(f.documents).as_posix(): 95,
                                        incoming.relative_to(f.documents).as_posix(): 40})
        return f, decisions

    def test_source_policy_and_csv_xlsx_threshold_are_supported(self):
        for suffix in ("csv", "xlsx"):
            with self.subTest(suffix=suffix), tempfile.TemporaryDirectory() as temp:
                f = Fixture(temp)
                p = f.file("Worker/Bulk/Batch 01/Other - Unknown.pdf")
                self.assertEqual(f.prepare(p, 80, suffix)["candidates"], [])
                self.assertEqual(f.queue["policy"]["batch_size"], 30)

    def test_ranking_requires_every_peer_not_just_the_new_document(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            decisions["peer_reviews"] = decisions["peer_reviews"][:1]
            with self.assertRaisesRegex(review.ReviewError, "Every affected category peer"):
                f.plan(decisions)

    def test_missing_quality_is_not_defaulted_to_zero(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            decisions["peer_reviews"][0].pop("score")
            with self.assertRaisesRegex(review.ReviewError, "real 0–100"):
                f.plan(decisions)

    def test_all_peers_rank_worst_to_best_and_route_by_type(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            targets = {op["source"]: op["target"] for op in plan["operations"]}
            self.assertEqual(targets["Worker/Overwrite Documents/DBS Document.pdf"],
                             "Worker/Overwrite Documents/DBS Document (01).pdf")
            self.assertEqual(targets["Worker/Bulk/Batch 01/Other - Unknown.pdf"],
                             "Worker/Overwrite Documents/DBS Document.pdf")
            self.assertEqual(len(plan["required_peers"]), 2)

    def test_actual_dated_policy_puts_newer_above_quality(self):
        with tempfile.TemporaryDirectory() as temp:
            f = Fixture(temp)
            older = f.file("Worker/Overwrite Documents/Certificate of Sponsorship - (01-01-2024).pdf")
            new = f.file("Worker/Bulk/Batch 01/Other - Unknown.pdf")
            f.prepare(new)
            decisions = f.decisions("Certificate of Sponsorship")
            for row in decisions["peer_reviews"]:
                row.update(date="2024-01-01" if row["path"] == older.relative_to(f.documents).as_posix() else "2025-01-01",
                           score=100 if row["path"] == older.relative_to(f.documents).as_posix() else 1)
            plan = f.plan(decisions)
            self.assertTrue(any(op["target"].endswith("Certificate of Sponsorship - (01-01-2025) (01).pdf") for op in plan["operations"]))

    def test_contract_signed_preference_matches_stage2(self):
        with tempfile.TemporaryDirectory() as temp:
            f = Fixture(temp)
            signed = f.file("Worker/Overwrite Documents/Employment Contract.pdf")
            incoming = f.file("Worker/Bulk/Batch 01/Other - Unknown.pdf")
            f.prepare(incoming)
            decisions = f.decisions("Employment Contract")
            for row in decisions["peer_reviews"]:
                row.update(signed=row["path"] == signed.relative_to(f.documents).as_posix(),
                           score=1 if row["path"] == signed.relative_to(f.documents).as_posix() else 100)
            plan = f.plan(decisions)
            self.assertTrue(any(op["source"] == signed.relative_to(f.documents).as_posix() and
                                op["target"].endswith("Employment Contract (01).pdf") for op in plan["operations"]))

    def test_other_collision_uses_free_number_without_scoring(self):
        with tempfile.TemporaryDirectory() as temp:
            f = Fixture(temp)
            f.file("Worker/Bulk/Batch 01/Other - payslip.pdf")
            incoming = f.file("Worker/Bulk/Batch 01/Other - Unknown.pdf")
            f.prepare(incoming)
            decisions = f.decisions("Other - payslip")
            decisions["peer_reviews"] = []
            plan = f.plan(decisions)
            self.assertEqual(plan["operations"][0]["target"], "Worker/Bulk/Batch 01/Other - payslip (01).pdf")

    def test_full_bulk_batch_uses_next_batch_without_relocating_existing_files(self):
        with tempfile.TemporaryDirectory() as temp:
            f = Fixture(temp)
            for index in range(30):
                f.file(f"Worker/Bulk/Batch 01/Other - item {index}.pdf")
            incoming = f.file("Worker/Overwrite Documents/Other - Unknown.pdf")
            f.prepare(incoming)
            plan = f.plan(f.decisions("Other - payslip"))
            self.assertEqual(len(plan["operations"]), 1)
            self.assertEqual(plan["operations"][0]["target"], "Worker/Bulk/Batch 02/Other - payslip.pdf")

    def test_new_peer_after_plan_blocks_before_any_rename(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            f.plan(decisions)
            f.file("Worker/Overwrite Documents/DBS Document (09).pdf")
            with self.assertRaisesRegex(review.ReviewError, "Worker documents changed"):
                review.apply_plan(f.plan_path, authorized=True)
            self.assertFalse((f.request / "REVIEW_TRANSACTION.json").exists())

    def test_apply_requires_authorization_and_preserves_every_hash_with_backups(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            with self.assertRaises(review.ReviewError):
                review.apply_plan(f.plan_path)
            result = review.apply_plan(f.plan_path, authorized=True)
            self.assertEqual(result["status"], "complete")
            for op in plan["operations"]:
                self.assertEqual(review.digest(f.documents / op["target"]), op["sha256"])
                self.assertEqual(review.digest(Path(result["backup_root"]) / op["source"]), op["sha256"])
            self.assertEqual(review.apply_plan(f.plan_path, authorized=True)["status"], "complete")

    def test_second_final_rename_failure_rolls_all_peers_back(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            original = review.move_no_replace
            calls = []
            def fail_once(source, target):
                calls.append((source, target))
                if len(calls) == 4:
                    raise OSError("synthetic final move failure")
                return original(source, target)
            with patch.object(review, "move_no_replace", side_effect=fail_once):
                with self.assertRaisesRegex(review.ReviewError, "rolled_back"):
                    review.apply_plan(f.plan_path, authorized=True)
            for op in plan["operations"]:
                self.assertEqual(review.digest(f.documents / op["source"]), op["sha256"])

    def test_crash_after_park_can_be_reconciled_and_rolled_back(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            original = review.move_no_replace
            calls = []
            def crash(source, target):
                calls.append((source, target))
                if len(calls) == 3:
                    raise SystemExit("synthetic process stop")
                return original(source, target)
            with patch.object(review, "move_no_replace", side_effect=crash):
                with self.assertRaises(SystemExit):
                    review.apply_plan(f.plan_path, authorized=True)
            result = review.recover_rollback(f.plan_path, authorized=True)
            self.assertEqual(result["status"], "rolled_back")
            for op in plan["operations"]:
                self.assertEqual(review.digest(f.documents / op["source"]), op["sha256"])

    def test_tampered_semantic_plan_rejected_even_with_recomputed_digest(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            plan["operations"][0]["target"] = "Worker/Overwrite Documents/BRP.pdf"
            plan.pop("plan_id")
            plan["plan_id"] = review.identity(json.dumps(plan, sort_keys=True))
            review.write_json(f.plan_path, plan)
            with self.assertRaisesRegex(review.ReviewError, "do not match"):
                review.apply_plan(f.plan_path, authorized=True)

    def test_records_include_collateral_peer_and_sync_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            f.plan(decisions)
            review.apply_plan(f.plan_path, authorized=True)
            ledger = f.base / "ledger"
            old = f.base / "legacy" / "Misnaming Record.xlsx"
            first = review.sync_records(f.request, ledger, legacy_record=old)
            second = review.sync_records(f.request, ledger, legacy_record=old)
            self.assertEqual(first["records"], 2)
            self.assertEqual(first, second)
            self.assertEqual(len(review.read_records(ledger / "review_records.jsonl")), 2)
            from openpyxl import load_workbook
            book = load_workbook(first["ledger"])
            self.assertEqual(book["Review Log"].max_row, 3)
            self.assertEqual(book["Run Log"].max_row, 2)
            self.assertEqual([c.value for c in book["Review Log"][1]], review.REVIEW_COLUMNS)
            self.assertEqual(book["Review Log"].freeze_panes, "A2")
            book.close()
            book = load_workbook(old)
            self.assertEqual(book.active.max_row, 3)
            book.close()

    def test_writer_lock_blocks_apply_before_document_mutation(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            f.plan(decisions)
            with review.writer_lock(f.documents):
                with self.assertRaisesRegex(review.ReviewError, "writer lock"):
                    review.apply_plan(f.plan_path, authorized=True)
            self.assertFalse((f.request / "REVIEW_TRANSACTION.json").exists())

    def test_unrelated_csv_is_rejected_instead_of_empty_success(self):
        with tempfile.TemporaryDirectory() as temp:
            audit = Path(temp) / "unrelated.csv"
            audit.write_text("Name,Value\na,1\n", encoding="utf-8")
            with self.assertRaisesRegex(review.ReviewError, "Audit headers"):
                review.read_audit(audit)

    def test_implemented_learning_requires_tests_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            f.plan(decisions)
            review.apply_plan(f.plan_path, authorized=True)
            ledger = f.base / "ledger"
            result = review.sync_records(f.request, ledger, ledger_path=ledger / "Custom Master.xlsx")
            entry = review.read_records(ledger / "review_records.jsonl")[0]["entry_id"]
            changes = {"reviewer": "Offline learning reviewer", "updates": [{"entry_id": entry,
                "expected_status": "Pending software review", "status": "Implemented",
                "version_or_commit": "synthetic-commit", "implementation_notes": "General rule verified"}]}
            path = f.request / "learning_updates.json"
            review.write_json(path, changes)
            with self.assertRaisesRegex(review.ReviewError, "regression-test evidence"):
                review.update_learning(path, ledger, ledger_path=result["ledger"], authorized=True)
            changes["updates"][0]["test_evidence"] = "Synthetic positive and negative controls passed"
            review.write_json(path, changes)
            first = review.update_learning(path, ledger, ledger_path=result["ledger"], authorized=True)
            second = review.update_learning(path, ledger, ledger_path=result["ledger"], authorized=True)
            self.assertEqual((first["updates"], second["updates"]), (1, 0))
            review.sync_records(f.request, ledger, ledger_path=result["ledger"])
            from openpyxl import load_workbook
            book = load_workbook(result["ledger"])
            self.assertEqual(book["Review Log"]["W2"].value, "Implemented")
            self.assertEqual(book["Review Log"]["X2"].value, "synthetic-commit")
            book.close()

    def test_learning_journal_first_failure_can_retry_without_duplicate_event(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            f.plan(decisions)
            review.apply_plan(f.plan_path, authorized=True)
            ledger = f.base / "ledger"
            result = review.sync_records(f.request, ledger)
            entry = review.read_records(ledger / "review_records.jsonl")[0]["entry_id"]
            path = f.request / "learning_updates.json"
            review.write_json(path, {"reviewer": "Offline reviewer", "updates": [{"entry_id": entry,
                "expected_status": "Pending software review", "status": "No software change",
                "implementation_notes": "Evidence does not justify a general change"}]})
            with patch.object(review, "_save_workbook_atomic", side_effect=OSError("synthetic locked workbook")):
                with self.assertRaises(OSError):
                    review.update_learning(path, ledger, authorized=True)
            self.assertEqual(review.update_learning(path, ledger, authorized=True)["updates"], 0)
            events = [r for r in review.read_records(ledger / "review_records.jsonl") if r["record_type"] == "improvement_update"]
            self.assertEqual(len(events), 1)
            from openpyxl import load_workbook
            book = load_workbook(result["ledger"])
            self.assertEqual(book["Review Log"]["W2"].value, "No software change")
            book.close()

    def test_crash_after_move_before_location_save_is_recoverable(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            original = review.move_no_replace
            calls = []
            def crash_after_move(source, target):
                calls.append((source, target))
                original(source, target)
                if len(calls) == 3:
                    raise SystemExit("synthetic crash before location persistence")
            with patch.object(review, "move_no_replace", side_effect=crash_after_move):
                with self.assertRaises(SystemExit):
                    review.apply_plan(f.plan_path, authorized=True)
            result = review.recover_rollback(f.plan_path, authorized=True)
            self.assertEqual(result["status"], "rolled_back")
            for op in plan["operations"]:
                self.assertEqual(review.digest(f.documents / op["source"]), op["sha256"])
            self.assertTrue(all(row["decision"] == "Error" for row in review.review_outcomes(plan, result)))

    def test_crash_mid_automatic_rollback_is_recoverable(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            plan = f.plan(decisions)
            original = review.move_no_replace
            calls = []
            def fail_then_crash(source, target):
                calls.append((source, target))
                if len(calls) == 4:
                    raise OSError("synthetic second final failure")
                original(source, target)
                if len(calls) == 5:
                    raise SystemExit("synthetic crash during rollback parking")
            with patch.object(review, "move_no_replace", side_effect=fail_then_crash):
                with self.assertRaises(SystemExit):
                    review.apply_plan(f.plan_path, authorized=True)
            result = review.recover_rollback(f.plan_path, authorized=True)
            self.assertEqual(result["status"], "rolled_back")
            for op in plan["operations"]:
                self.assertEqual(review.digest(f.documents / op["source"]), op["sha256"])

    def test_review_only_proposals_and_keeps_sync_without_document_permission(self):
        for choice in ("rename", "keep", "defer"):
            with self.subTest(choice=choice), tempfile.TemporaryDirectory() as temp:
                f, decisions = self.ranking_fixture(temp)
                decisions["decisions"][0]["decision"] = choice
                before = {row["path"]: row["sha256"] for row in f.queue["inventory"]}
                f.plan(decisions)
                result = review.finalize_review(f.plan_path)
                self.assertEqual(result["status"], "review_only")
                review.sync_records(f.request, f.base / "ledger")
                self.assertFalse((f.documents / ".docreview_batch_writer.lock").exists())
                self.assertFalse((f.request / "backups").exists())
                for path, sha in before.items():
                    self.assertEqual(review.digest(f.documents / path), sha)
                records = review.read_records(f.base / "ledger" / "review_records.jsonl")
                self.assertTrue(all(not row["rename_applied"] and row["applied_relative_path"] is None for row in records))
                if choice == "rename":
                    self.assertTrue(all(row["apply_outcome"] == "proposed_not_applied" for row in records))

    def test_new_review_attempt_can_record_later_correction_of_same_deferred_case(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            decisions["decisions"][0]["decision"] = "defer"
            f.plan(decisions)
            review.finalize_review(f.plan_path)
            ledger = f.base / "ledger"
            review.sync_records(f.request, ledger)
            first_candidate = f.queue["candidates"][0]
            first_run = f.queue["run_id"]
            f.request = f.base / "second-request"
            f.request.mkdir()
            f.queue_path = f.request / "review_queue.json"
            f.queue = review.prepare(f.audit, "Synthetic Home", f.documents, f.source, f.queue_path)
            second_candidate = f.queue["candidates"][0]
            self.assertNotEqual(first_run, f.queue["run_id"])
            self.assertNotEqual(first_candidate["candidate_id"], second_candidate["candidate_id"])
            self.assertEqual(first_candidate["case_id"], second_candidate["case_id"])
            f.plan(f.decisions())
            review.apply_plan(f.plan_path, authorized=True)
            review.sync_records(f.request, ledger)
            review.sync_records(f.request, ledger)
            rows = review.read_records(ledger / "review_records.jsonl")
            case = [row for row in rows if row.get("case_id") == first_candidate["case_id"]]
            self.assertEqual([row["decision"] for row in case], ["Defer", "Rename"])
            self.assertTrue(case[-1]["rename_applied"])

    def test_duplicate_request_writer_blocks_and_saved_plan_is_immutable(self):
        with tempfile.TemporaryDirectory() as temp:
            f, decisions = self.ranking_fixture(temp)
            f.plan(decisions)
            with self.assertRaisesRegex(review.ReviewError, "Plan already exists"):
                f.plan(decisions)
            with review.writer_lock(f.request, ".ai_review_request.lock"):
                with self.assertRaisesRegex(review.ReviewError, "writer lock"):
                    review.finalize_review(f.plan_path)
            self.assertFalse((f.request / "REVIEW_TRANSACTION.json").exists())


if __name__ == "__main__":
    unittest.main()
